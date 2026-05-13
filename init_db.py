"""
init_db.py — Khởi tạo Database từ file CSV hoặc XLSX
Tạo schema + import dữ liệu nguyên liệu vào PostgreSQL.

Usage:
    python init_db.py                           # Interactive mode
    python init_db.py --file data/materials.csv  # Direct import
    python init_db.py --file data/materials.xlsx --sheet Sheet1
    python init_db.py --reset                    # Xóa và tạo lại DB (⚠️ mất dữ liệu)
"""
import argparse
import csv
import os
import sys
import time
from pathlib import Path

# ── Setup path ────────────────────────────────────────────
ROOT_DIR = Path(__file__).parent
sys.path.insert(0, str(ROOT_DIR))

try:
    import psycopg2
    from psycopg2 import sql as psql
except ImportError:
    print("❌ psycopg2 chưa được cài đặt. Chạy: pip install psycopg2-binary")
    sys.exit(1)

from config import DB_CONFIG

# ── Constants ─────────────────────────────────────────────
REQUIRED_COLUMNS = {"ma_hc", "ten_nguyen_lieu", "ten_ke_toan"}
OPTIONAL_COLUMNS = {"chuc_nang"}
ALL_COLUMNS = REQUIRED_COLUMNS | OPTIONAL_COLUMNS

CREATE_MATERIALS_TABLE = """
CREATE TABLE IF NOT EXISTS raw_materials (
    id SERIAL PRIMARY KEY,
    ma_hc VARCHAR(50) NOT NULL,
    ten_nguyen_lieu VARCHAR(500) NOT NULL,
    chuc_nang VARCHAR(500) DEFAULT '',
    ten_ke_toan VARCHAR(500) DEFAULT '',
    created_at TIMESTAMP DEFAULT NOW(),
    UNIQUE(ma_hc)
);
"""

CREATE_ALIASES_TABLE = """
CREATE TABLE IF NOT EXISTS material_aliases (
    id SERIAL PRIMARY KEY,
    ocr_name VARCHAR(500) NOT NULL UNIQUE,
    matched_ma_hc VARCHAR(50) NOT NULL,
    matched_ten_nl VARCHAR(500),
    confidence VARCHAR(20) DEFAULT 'manual',
    created_at TIMESTAMP DEFAULT NOW()
);
"""

CREATE_INDEX = """
CREATE INDEX IF NOT EXISTS idx_materials_ma_hc ON raw_materials(ma_hc);
CREATE INDEX IF NOT EXISTS idx_materials_ten_nl ON raw_materials(ten_nguyen_lieu);
CREATE INDEX IF NOT EXISTS idx_aliases_ocr_name ON material_aliases(ocr_name);
"""


def wait_for_db(max_retries: int = 30, delay: float = 2.0) -> bool:
    """Chờ PostgreSQL sẵn sàng (hữu ích khi Docker vừa start)."""
    print("⏳ Đang chờ PostgreSQL sẵn sàng...", end="", flush=True)
    for i in range(max_retries):
        try:
            conn = psycopg2.connect(**DB_CONFIG)
            conn.close()
            print(" ✅")
            return True
        except psycopg2.OperationalError:
            print(".", end="", flush=True)
            time.sleep(delay)
    print(" ❌ Timeout!")
    return False


def create_schema(conn):
    """Tạo các table cần thiết."""
    with conn.cursor() as cur:
        cur.execute(CREATE_MATERIALS_TABLE)
        cur.execute(CREATE_ALIASES_TABLE)
        cur.execute(CREATE_INDEX)
    conn.commit()
    print("✅ Schema đã được tạo (raw_materials + material_aliases)")


def reset_database(conn):
    """⚠️ Xóa toàn bộ dữ liệu và tạo lại schema."""
    print("⚠️  CẢNH BÁO: Thao tác này sẽ XÓA TOÀN BỘ dữ liệu!")
    confirm = input("   Nhập 'YES' để xác nhận: ").strip()
    if confirm != "YES":
        print("   ❌ Đã hủy.")
        return False

    with conn.cursor() as cur:
        cur.execute("DROP TABLE IF EXISTS material_aliases CASCADE;")
        cur.execute("DROP TABLE IF EXISTS raw_materials CASCADE;")
    conn.commit()
    print("🗑️  Đã xóa các table cũ.")
    create_schema(conn)
    return True


def read_csv_file(filepath: str) -> list[dict]:
    """Đọc file CSV → list[dict]."""
    records = []
    encodings = ["utf-8-sig", "utf-8", "latin-1", "cp1252"]

    for encoding in encodings:
        try:
            with open(filepath, "r", encoding=encoding) as f:
                reader = csv.DictReader(f)
                headers = set(reader.fieldnames or [])

                # Validate required columns
                missing = REQUIRED_COLUMNS - headers
                if missing:
                    print(f"❌ File thiếu các cột bắt buộc: {missing}")
                    print(f"   Các cột có trong file: {headers}")
                    print(f"   Các cột cần thiết: {REQUIRED_COLUMNS}")
                    return []

                for row in reader:
                    record = {
                        "ma_hc": (row.get("ma_hc") or "").strip(),
                        "ten_nguyen_lieu": (row.get("ten_nguyen_lieu") or "").strip(),
                        "chuc_nang": (row.get("chuc_nang") or "").strip(),
                        "ten_ke_toan": (row.get("ten_ke_toan") or "").strip(),
                    }
                    # Skip empty rows
                    if record["ma_hc"] and record["ten_nguyen_lieu"]:
                        records.append(record)

            print(f"📄 Đọc {len(records)} records từ CSV (encoding: {encoding})")
            return records
        except UnicodeDecodeError:
            continue
        except Exception as e:
            print(f"❌ Lỗi đọc CSV: {e}")
            return []

    print("❌ Không thể đọc file CSV với bất kỳ encoding nào")
    return []


def read_xlsx_file(filepath: str, sheet_name: str = None) -> list[dict]:
    """Đọc file XLSX → list[dict]."""
    try:
        import openpyxl
    except ImportError:
        print("❌ openpyxl chưa được cài đặt. Chạy: pip install openpyxl")
        return []

    try:
        wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)

        # Chọn sheet
        if sheet_name and sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
        else:
            if sheet_name:
                print(f"⚠️  Sheet '{sheet_name}' không tìm thấy. Dùng sheet đầu tiên.")
            ws = wb.active
            print(f"📊 Sử dụng sheet: '{ws.title}'")

        # Đọc header từ dòng đầu tiên
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            print("❌ File rỗng!")
            return []

        headers = [str(h).strip().lower() if h else "" for h in rows[0]]

        # Validate required columns
        header_set = set(headers)
        missing = REQUIRED_COLUMNS - header_set
        if missing:
            print(f"❌ File thiếu các cột bắt buộc: {missing}")
            print(f"   Các cột có trong file: {header_set}")
            print(f"   Các cột cần thiết: {REQUIRED_COLUMNS}")
            return []

        # Parse data rows
        records = []
        for row in rows[1:]:
            row_dict = {}
            for i, val in enumerate(row):
                if i < len(headers) and headers[i] in ALL_COLUMNS:
                    row_dict[headers[i]] = str(val).strip() if val else ""

            record = {
                "ma_hc": row_dict.get("ma_hc", "").strip(),
                "ten_nguyen_lieu": row_dict.get("ten_nguyen_lieu", "").strip(),
                "chuc_nang": row_dict.get("chuc_nang", "").strip(),
                "ten_ke_toan": row_dict.get("ten_ke_toan", "").strip(),
            }
            if record["ma_hc"] and record["ten_nguyen_lieu"]:
                records.append(record)

        wb.close()
        print(f"📄 Đọc {len(records)} records từ XLSX")
        return records

    except Exception as e:
        print(f"❌ Lỗi đọc XLSX: {e}")
        return []


def import_data(conn, records: list[dict]) -> tuple[int, int]:
    """Import records vào raw_materials. Returns (inserted, skipped)."""
    inserted = 0
    skipped = 0

    with conn.cursor() as cur:
        for rec in records:
            try:
                cur.execute("""
                    INSERT INTO raw_materials (ma_hc, ten_nguyen_lieu, chuc_nang, ten_ke_toan)
                    VALUES (%s, %s, %s, %s)
                    ON CONFLICT (ma_hc) DO UPDATE
                    SET ten_nguyen_lieu = EXCLUDED.ten_nguyen_lieu,
                        chuc_nang = EXCLUDED.chuc_nang,
                        ten_ke_toan = EXCLUDED.ten_ke_toan
                """, (
                    rec["ma_hc"],
                    rec["ten_nguyen_lieu"],
                    rec["chuc_nang"],
                    rec["ten_ke_toan"],
                ))
                inserted += 1
            except Exception as e:
                print(f"  ⚠️ Skip [{rec['ma_hc']}]: {e}")
                skipped += 1
                conn.rollback()

    conn.commit()
    return inserted, skipped


def get_table_stats(conn) -> dict:
    """Lấy thống kê hiện tại của DB."""
    stats = {}
    with conn.cursor() as cur:
        try:
            cur.execute("SELECT COUNT(*) FROM raw_materials")
            stats["materials"] = cur.fetchone()[0]
        except Exception:
            stats["materials"] = 0
            conn.rollback()

        try:
            cur.execute("SELECT COUNT(*) FROM material_aliases")
            stats["aliases"] = cur.fetchone()[0]
        except Exception:
            stats["aliases"] = 0
            conn.rollback()

    return stats


def interactive_mode():
    """Chế độ tương tác — hỏi user chọn file."""
    print("\n" + "=" * 60)
    print("📦 KHỞI TẠO DATABASE — Chế độ tương tác")
    print("=" * 60)

    # Tìm file trong thư mục data/
    data_dir = ROOT_DIR / "data"
    available_files = []
    if data_dir.exists():
        for f in data_dir.iterdir():
            if f.suffix.lower() in (".csv", ".xlsx"):
                available_files.append(f)

    if available_files:
        print(f"\n📂 Tìm thấy {len(available_files)} file trong thư mục data/:")
        for i, f in enumerate(available_files, 1):
            print(f"   {i}. {f.name} ({f.stat().st_size / 1024:.1f} KB)")

        print(f"   {len(available_files) + 1}. Nhập đường dẫn khác")
        choice = input(f"\n🔹 Chọn (1-{len(available_files) + 1}): ").strip()

        try:
            idx = int(choice) - 1
            if 0 <= idx < len(available_files):
                return str(available_files[idx])
        except ValueError:
            pass

    filepath = input("\n📁 Nhập đường dẫn file CSV hoặc XLSX: ").strip()
    filepath = filepath.strip('"').strip("'")  # Bỏ quotes nếu user kéo thả file
    return filepath


def main():
    parser = argparse.ArgumentParser(
        description="Khởi tạo Database từ file CSV/XLSX",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Ví dụ:
  python init_db.py                              # Interactive mode
  python init_db.py --file data/materials.csv    # Import từ CSV
  python init_db.py --file data/materials.xlsx   # Import từ XLSX
  python init_db.py --file data/materials.xlsx --sheet "Sheet1"
  python init_db.py --schema-only                # Chỉ tạo schema, không import
  python init_db.py --reset                      # Xóa và tạo lại DB

Định dạng file yêu cầu:
  Cột bắt buộc: ma_hc, ten_nguyen_lieu, ten_ke_toan
  Cột tùy chọn: chuc_nang
        """,
    )
    parser.add_argument("--file", "-f", help="Đường dẫn file CSV hoặc XLSX")
    parser.add_argument("--sheet", "-s", help="Tên sheet (chỉ cho XLSX)", default=None)
    parser.add_argument("--schema-only", action="store_true", help="Chỉ tạo schema, không import data")
    parser.add_argument("--reset", action="store_true", help="Xóa và tạo lại toàn bộ DB")
    args = parser.parse_args()

    print("\n" + "=" * 60)
    print("🏗️  MATERIAL OCR — DATABASE INITIALIZATION")
    print("=" * 60)
    print(f"   Host: {DB_CONFIG['host']}:{DB_CONFIG['port']}")
    print(f"   Database: {DB_CONFIG['database']}")
    print(f"   User: {DB_CONFIG['user']}")
    print("=" * 60)

    # ── Step 1: Chờ DB sẵn sàng ──
    if not wait_for_db():
        print("\n❌ Không thể kết nối PostgreSQL!")
        print("   Hãy chắc chắn Docker đang chạy:")
        print("   > docker-compose up -d")
        sys.exit(1)

    conn = psycopg2.connect(**DB_CONFIG)

    try:
        # ── Step 2: Reset nếu cần ──
        if args.reset:
            if not reset_database(conn):
                sys.exit(0)
        else:
            create_schema(conn)

        # ── Step 3: Schema only mode ──
        if args.schema_only:
            stats = get_table_stats(conn)
            print(f"\n📊 DB Stats: {stats['materials']} nguyên liệu, {stats['aliases']} aliases")
            print("✅ Schema đã sẵn sàng. Không import data (--schema-only).")
            return

        # ── Step 4: Xác định file nguồn ──
        filepath = args.file
        if not filepath:
            filepath = interactive_mode()

        if not filepath:
            print("❌ Không có file nào được chọn.")
            sys.exit(1)

        filepath = os.path.abspath(filepath)
        if not os.path.exists(filepath):
            print(f"❌ File không tồn tại: {filepath}")
            sys.exit(1)

        print(f"\n📂 File: {filepath}")

        # ── Step 5: Đọc file ──
        ext = os.path.splitext(filepath)[1].lower()
        if ext == ".csv":
            records = read_csv_file(filepath)
        elif ext in (".xlsx", ".xls"):
            records = read_xlsx_file(filepath, args.sheet)
        else:
            print(f"❌ Định dạng không được hỗ trợ: {ext}")
            print("   Chỉ hỗ trợ: .csv, .xlsx")
            sys.exit(1)

        if not records:
            print("❌ Không có dữ liệu để import.")
            sys.exit(1)

        # ── Step 6: Import ──
        print(f"\n📥 Đang import {len(records)} records...")
        inserted, skipped = import_data(conn, records)

        # ── Step 7: Kết quả ──
        stats = get_table_stats(conn)
        print("\n" + "=" * 60)
        print("✅ HOÀN TẤT!")
        print(f"   📥 Import: {inserted} thành công, {skipped} bỏ qua")
        print(f"   📊 DB tổng: {stats['materials']} nguyên liệu, {stats['aliases']} aliases")
        print("=" * 60)

    finally:
        conn.close()


if __name__ == "__main__":
    main()
