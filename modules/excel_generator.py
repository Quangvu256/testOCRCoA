"""
excel_generator.py — Xuất file .xlsx với conditional formatting
HSD < 90 ngày → bôi đỏ
"""
import os
import sys
from datetime import datetime
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from config import EXPIRY_WARNING_DAYS
from models import OutputRecord

# ── Styles ────────────────────────────────────────────────
RED_FILL = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
WHITE_FONT = Font(color="FFFFFF", bold=True)
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
BODY_FONT = Font(size=11)
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
CENTER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT_ALIGN = Alignment(horizontal="left", vertical="center", wrap_text=True)

# ── Cấu trúc cột ─────────────────────────────────────────
COLUMNS = [
    ("STT", 6),
    ("NgayNhap", 14),
    ("MaHC", 10),
    ("TenHC", 35),
    ("TenHC(Ke Toan)", 30),
    ("SoLo", 18),
    ("SoLuong", 14),
    ("SoPack", 12),
    ("NgaySX", 14),
    ("HSD", 14),
    ("TenCongTy", 28),
    ("Titrong", 16),
    ("Tile", 16),
    ("Trongluongrieng", 20),
    ("HinhDang", 24),
    ("HSDConLaiNgay", 18),
]


def generate_xlsx(records: list[OutputRecord], output_dir: str = "outputs") -> str:
    """
    Tạo file .xlsx từ danh sách OutputRecord.
    
    Conditional formatting:
    - HSD < 90 ngày → bôi đỏ ô HSD + HSD còn lại
    
    Returns: đường dẫn file đã tạo.
    """
    # Đảm bảo output dir tồn tại
    os.makedirs(output_dir, exist_ok=True)

    wb = Workbook()
    ws = wb.active

    # Sheet name: "Nhập NL DD-MM-YYYY"
    today_str = datetime.now().strftime("%d-%m-%Y")
    ws.title = f"Nhập NL {today_str}"

    # ── Header row ────────────────────────────────────────
    for col_idx, (col_name, col_width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = THIN_BORDER
        cell.alignment = CENTER_ALIGN
        ws.column_dimensions[get_column_letter(col_idx)].width = col_width

    # Freeze header
    ws.freeze_panes = "A2"

    # ── Data rows ─────────────────────────────────────────
    for row_idx, record in enumerate(records, start=2):
        # Tính HSD còn lại
        record.calculate_remaining_days()

        row_data = [
            record.stt,
            record.ngay_nhap.strftime("%d/%m/%Y") if record.ngay_nhap else "",
            record.ma_hc,
            record.ten_nl,
            record.ten_ke_toan,
            record.so_lo,
            record.so_luong_kg,
            record.so_pack,
            record.ngay_sx.strftime("%d/%m/%Y") if record.ngay_sx else "",
            record.hsd.strftime("%d/%m/%Y") if record.hsd else "",
            record.ten_cong_ty or "",
            record.ti_trong or "",
            record.ti_le or "",
            record.trong_luong_rieng or "",
            record.hinh_dang or "",
            record.thoi_gian_hsd_con_lai,
        ]

        for col_idx, value in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = BODY_FONT
            cell.border = THIN_BORDER
            cell.alignment = CENTER_ALIGN if col_idx in (1, 2, 3, 6, 7, 8, 9, 10, 12, 13, 14, 16) else LEFT_ALIGN

        # ── Conditional formatting: HSD < 90 ngày → đỏ ──
        if record.thoi_gian_hsd_con_lai is not None and record.thoi_gian_hsd_con_lai < EXPIRY_WARNING_DAYS:
            # Cột HSD (index 10)
            hsd_cell = ws.cell(row=row_idx, column=10)
            hsd_cell.fill = RED_FILL
            hsd_cell.font = WHITE_FONT

            # Cột HSD còn lại (index 16)
            remaining_cell = ws.cell(row=row_idx, column=16)
            remaining_cell.fill = RED_FILL
            remaining_cell.font = WHITE_FONT

            # Thêm text mô tả
            if record.thoi_gian_hsd_con_lai is not None:
                remaining_cell.value = f"{record.thoi_gian_hsd_con_lai} ngày"

    # ── Save file ─────────────────────────────────────────
    filename = f"nhap_nl_{today_str}.xlsx"
    filepath = os.path.join(output_dir, filename)
    wb.save(filepath)

    return filepath
